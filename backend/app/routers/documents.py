from fastapi import APIRouter, Depends, HTTPException, UploadFile, File
from sqlalchemy.orm import Session
from app.database import get_db
from app import models, schemas
from app.auth import get_current_user
from app.services.vectorstore import store_document, delete_document_chunks
from docx import Document as DocxDocument
from pptx import Presentation
import fitz
import uuid
import os

router = APIRouter()

UPLOAD_DIR = "uploads"

MAX_FILE_SIZE = 50 * 1024 * 1024
MAX_FILE_SIZE_MB = MAX_FILE_SIZE // (1024 * 1024)
os.makedirs(UPLOAD_DIR, exist_ok=True)

def get_original_upload_paths(doc: models.Document) -> list[str]:
    if doc.file_path:
        return [doc.file_path]

    suffix = f"_{doc.filename}"
    matches = [
        os.path.join(UPLOAD_DIR, name)
        for name in os.listdir(UPLOAD_DIR)
        if name.endswith(suffix)
    ]
    return matches if len(matches) == 1 else []

def extract_text(file_path: str, filename: str) -> str:
    ext = filename.lower().split(".")[-1]

    if ext == "pdf":
        doc = fitz.open(file_path)
        return "\n".join(page.get_text() for page in doc)

    elif ext in ["txt", "md"]:
        with open(file_path, "r", encoding="utf-8") as f:
            return f.read()

    elif ext == "docx":
        doc = DocxDocument(file_path)
        return "\n".join(
            para.text for para in doc.paragraphs if para.text.strip()
        )

    elif ext == "pptx":
        prs = Presentation(file_path)
        text_parts = []
        for slide in prs.slides:
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text.strip():
                    text_parts.append(shape.text)
        return "\n".join(text_parts)

    else:
        raise HTTPException(
            status_code=400,
            detail="Unsupported file type. Use PDF, TXT, MD, DOCX or PPTX."
        )


@router.post("/upload", response_model=schemas.DocumentOut)
def upload_document(
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    current_user: models.User = Depends(get_current_user)
):
    allowed = ["pdf", "txt", "md", "docx", "pptx"]
    ext = file.filename.lower().split(".")[-1]
    if ext not in allowed:
        raise HTTPException(
            status_code=400,
            detail="Only PDF, TXT, MD, DOCX and PPTX files are supported."
        )

    contents = file.file.read()
    if len(contents) > MAX_FILE_SIZE:
        raise HTTPException(
            status_code=400,
            detail=f"File too large. Maximum size is {MAX_FILE_SIZE_MB}MB."
        )

    existing = db.query(models.Document).filter(
        models.Document.user_id == current_user.id,
        models.Document.filename == file.filename
    ).first()
    if existing:
        raise HTTPException(
            status_code=400,
            detail="A document with this name already exists."
        )

    file_id = str(uuid.uuid4())
    file_path = os.path.join(UPLOAD_DIR, f"{file_id}_{file.filename}")

    with open(file_path, "wb") as f:
        f.write(contents)

    try:
        text = extract_text(file_path, file.filename)
    except HTTPException:
        os.remove(file_path)
        raise
    except Exception as e:
        os.remove(file_path)
        raise HTTPException(
            status_code=500,
            detail=f"Failed to extract text: {str(e)}"
        )

    if not text.strip():
        os.remove(file_path)
        raise HTTPException(
            status_code=400,
            detail="Document appears to be empty."
        )

    doc_record = models.Document(
        id=uuid.uuid4(),
        user_id=current_user.id,
        filename=file.filename,
        file_path=file_path
    )
    db.add(doc_record)
    db.commit()
    db.refresh(doc_record)

    text_path = os.path.join(UPLOAD_DIR, f"{str(doc_record.id)}.txt")
    with open(text_path, "w", encoding="utf-8") as f:
        f.write(text)

    try:
        chunk_count = store_document(
            doc_id=str(doc_record.id),
            user_id=str(current_user.id),
            text=text,
            chunk_size=current_user.chunk_size
        )
        doc_record.status = "ready"
        doc_record.chunk_count = chunk_count
    except Exception as e:
        doc_record.status = "failed"
        doc_record.chunk_count = 0

    db.commit()
    db.refresh(doc_record)

    return doc_record


@router.get("/", response_model=list[schemas.DocumentOut])
def list_documents(
    db: Session = Depends(get_db),
    current_user: models.User = Depends(get_current_user)
):
    return db.query(models.Document).filter(
        models.Document.user_id == current_user.id
    ).all()


@router.delete("/{document_id}")
def delete_document(
    document_id: str,
    db: Session = Depends(get_db),
    current_user: models.User = Depends(get_current_user)
):
    doc = db.query(models.Document).filter(
        models.Document.id == document_id,
        models.Document.user_id == current_user.id
    ).first()

    if not doc:
        raise HTTPException(status_code=404, detail="Document not found.")

    for path in [
        *get_original_upload_paths(doc),
        os.path.join(UPLOAD_DIR, f"{document_id}.txt"),
    ]:
        if os.path.exists(path):
            os.remove(path)

    delete_document_chunks(document_id)
    db.delete(doc)
    db.commit()
    return {"message": "Document deleted successfully"}
